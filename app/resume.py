from docx import Document
from app.config import MASTER_RESUME,GENERATED_DIR
def read_master_resume():return '\n'.join(p.text.strip() for p in Document(MASTER_RESUME).paragraphs if p.text.strip())
def save_resume(job_id,d,variant):
    GENERATED_DIR.mkdir(parents=True,exist_ok=True);doc=Document();doc.add_heading('SAI JAGAN YALAMANCHILI',0);doc.add_paragraph(d.get('headline','DATA ANALYST'));doc.add_paragraph('Auburn Hills, MI | +1 989-933-2400 | saijagan.yalamanchili7@gmail.com | linkedin.com/in/sai-jagan-yalamanchili-24b247355');doc.add_heading('PROFESSIONAL SUMMARY',1);doc.add_paragraph(d.get('summary',''));doc.add_heading('CORE SKILLS',1);doc.add_paragraph(' • '.join(d.get('skills',[])));doc.add_heading('PROFESSIONAL EXPERIENCE',1)
    for e in d.get('experience',[]):
        p=doc.add_paragraph();p.add_run(e.get('employer','')).bold=True;doc.add_paragraph(' | '.join(x for x in [e.get('title',''),e.get('dates','')] if x))
        for b in e.get('bullets',[]):doc.add_paragraph(b,style='List Bullet')
    path=GENERATED_DIR/f'job_{job_id}_Resume_{variant}.docx';doc.save(path);return str(path)
def save_cover(job_id,company,title,body,variant):
    doc=Document();doc.add_paragraph('Sai Jagan Yalamanchili');doc.add_paragraph('Auburn Hills, MI | saijagan.yalamanchili7@gmail.com | +1 989-933-2400');doc.add_paragraph();doc.add_paragraph(f'Re: {title} — {company}');doc.add_paragraph()
    for b in body.split('\n\n'):
        if b.strip():doc.add_paragraph(b.strip())
    path=GENERATED_DIR/f'job_{job_id}_Cover_Letter_{variant}.docx';doc.save(path);return str(path)
